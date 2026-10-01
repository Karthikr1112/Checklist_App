"""Import stores from a Plant Master Excel file.

Usage:
    python manage.py import_stores "Plant Master.xlsx"

Maps columns:
    Plant   -> Store.code
    Name    -> Store.name
    City    -> Store.city
"""
from __future__ import annotations

from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

from apps.stores.models import Store


class Command(BaseCommand):
    help = "Import stores from a Plant Master .xlsx file."

    def add_arguments(self, parser):
        parser.add_argument("path", help="Path to the Plant Master .xlsx file")
        parser.add_argument(
            "--sheet",
            default=None,
            help="Sheet name (defaults to the first sheet).",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Parse and print what would be imported without saving.",
        )

    def handle(self, *args, **options):
        try:
            import openpyxl  # noqa: WPS433
        except ImportError as exc:
            raise CommandError("openpyxl is required: pip install openpyxl") from exc

        path = Path(options["path"]).expanduser().resolve()
        if not path.exists():
            raise CommandError(f"File not found: {path}")

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[options["sheet"]] if options["sheet"] else wb[wb.sheetnames[0]]

        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(rows)]
        try:
            idx_code = header.index("Plant")
            idx_name = header.index("Name")
            idx_city = header.index("City")
        except ValueError as exc:
            raise CommandError(
                f"Required column missing in header {header!r}: {exc}"
            ) from exc

        created = updated = skipped = 0
        for row in rows:
            if not row or all(c is None for c in row):
                continue
            code = self._clean(row[idx_code])
            name = self._clean(row[idx_name])
            city = self._clean(row[idx_city])
            if not code or not name:
                skipped += 1
                continue

            defaults = {"name": name, "city": city, "is_active": True}
            if options["dry_run"]:
                self.stdout.write(f"[dry-run] {code} :: {name} ({city})")
                continue

            obj, was_created = Store.objects.update_or_create(code=code, defaults=defaults)
            if was_created:
                created += 1
            else:
                updated += 1

        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("Dry run — nothing saved."))
        else:
            self.stdout.write(
                self.style.SUCCESS(
                    f"Done. Created: {created}, Updated: {updated}, Skipped: {skipped}"
                )
            )

    @staticmethod
    def _clean(value) -> str:
        if value is None:
            return ""
        return str(value).strip()
